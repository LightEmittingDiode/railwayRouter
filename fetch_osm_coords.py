"""
fetch_osm_coords.py
순수 한글 역명만 정제하여 OpenStreetMap Nominatim API로 좌표를 조회하고 저장합니다.
"""
import urllib.request, urllib.parse, json, sys, time, csv, os, re
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

CACHE_FILE = r'C:\Users\harry\Desktop\osm_coords_cache.json'
CSV_PATH = r'C:\Users\harry\Downloads\한국철도공사_역 위치 정보_20240401 (2).csv'
FILES = [
    r'C:\Users\harry\Downloads\KTX_timetable.xlsx',
    r'C:\Users\harry\Downloads\ITXC_timetable.xlsx',
    r'C:\Users\harry\Downloads\Local_timetable.xlsx',
    r'C:\Users\harry\Downloads\SRT_timetable.xlsx',
]

S_PAT = re.compile(r'^S\d{4}$')
INVALID_EXACT = {
    '열차번호','열차종별','비 고','비고','종착역','시발역','열차',
    'KTX','SRT','무궁화','ITX','누리로','새마을','KTX-산천','KTX-청룡',
    'ITX-마음','ITX-새마을','KTX-이음','KTX_산천', '편성'  # '편성' 제외!
}

hangul_pattern = re.compile(r'[\uac00-\ud7a3]')
hanja_pattern = re.compile(r'[\u4e00-\u9fff\u3400-\u4dbf]')

def clean_station_name(v):
    if v is None: return None
    s = str(v).strip()
    if not s or s == '_': return None
    if S_PAT.match(s): return None
    if s in INVALID_EXACT: return None
    if hanja_pattern.search(s): return None
    
    # 통합 예외 처리
    if s in ('홍성(도착)', '홍성(출발)'):
        return '홍성'
        
    if not hangul_pattern.search(s): return None
    
    if re.search(r'선|발|행|예시|년|월|일|보기|주말|평일|운행|정차|통과|시표|시간표', s) and len(s) > 4:
        return None
    if re.search(r'^[월화수목금토일]+$', s):
        return None
        
    s_clean = re.sub(r'\s+', '', s)
    if len(s_clean) > 10 or len(s_clean) < 2: return None
    return s_clean

# ─── 1. CSV 좌표 로드 ───
csv_coords = {}
if os.path.exists(CSV_PATH):
    with open(CSV_PATH, 'r', encoding='euc-kr', errors='ignore') as f:
        for row in csv.DictReader(f):
            name = row.get('역명','').strip()
            name_clean = clean_station_name(name) or re.sub(r'\s+', '', name)
            try:
                lat, lon = float(row['위도']), float(row['경도'])
                if lat and lon:
                    csv_coords[name_clean] = (lat, lon)
            except: pass

# ─── 2. OSM 캐시 로드 ───
osm_cache = {}
if os.path.exists(CACHE_FILE):
    with open(CACHE_FILE, 'r', encoding='utf-8') as f:
        osm_cache = json.load(f)

# '편성' 등 불필요 키 삭제
if '편성' in osm_cache: del osm_cache['편성']

# 판교(경기) 수동 고정 좌표 (성남시 백현동 판교역)
osm_cache['판교(경기)'] = {'lat': 37.39472, 'lon': 127.11153, 'source': 'custom', 'display': '판교역 경기도 성남시'}

# ─── 3. 역 목록 정제 ───
all_stations = set()
for fpath in FILES:
    if not os.path.exists(fpath): continue
    wb = openpyxl.load_workbook(fpath, data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            for val in row:
                st = clean_station_name(val)
                if st: all_stations.add(st)

missing = [s for s in sorted(all_stations) if s not in csv_coords and s not in osm_cache]

print(f"📊 [역 목록 정제 완료]")
print(f"• 전체 한글 역 수: {len(all_stations)}개")
print(f"• CSV 기본 매칭: {sum(1 for s in all_stations if s in csv_coords)}개")
print(f"• OSM 조회 필요 (누락 역): {len(missing)}개\n")

# ─── 4. Nominatim API 조회 ───
def geocode(name):
    if name == '판교(경기)':
        return 37.39472, 127.11153, '판교역 경기도 성남시'
    
    search_q = name
    if '판교' in name and '경기' in name:
        search_q = '판교역 성남시 경기도'
        
    queries = [
        f'{search_q}역 대한민국',
        f'{search_q} 대한민국',
    ]
    for q in queries:
        try:
            url = f'https://nominatim.openstreetmap.org/search?q={urllib.parse.quote(q)}&format=json&limit=2&accept-language=ko&countrycodes=kr'
            req = urllib.request.Request(url, headers={'User-Agent': 'KorailRouteMap/1.0 (educational)'})
            with urllib.request.urlopen(req, timeout=6) as r:
                data = json.loads(r.read())
            if data:
                for item in data:
                    lat, lon = float(item['lat']), float(item['lon'])
                    if 33 <= lat <= 39 and 124 <= lon <= 132:
                        return lat, lon, item.get('display_name','')[:50]
        except Exception:
            pass
        time.sleep(0.3)
    return None

print(f"OSM 위치 조회 시작...")

try:
    for i, name in enumerate(missing):
        result = geocode(name)
        time.sleep(1.1)
        
        if result:
            lat, lon, display = result
            osm_cache[name] = {'lat': lat, 'lon': lon, 'source': 'osm', 'display': display}
            print(f"[{i+1}/{len(missing)}] ✅ {name}: {lat:.5f}, {lon:.5f} ({display})")
        else:
            osm_cache[name] = {'lat': None, 'lon': None, 'source': 'osm', 'display': ''}
            print(f"[{i+1}/{len(missing)}] ❌ {name}: 위치 미발견")
        
        with open(CACHE_FILE, 'w', encoding='utf-8') as f:
            json.dump(osm_cache, f, ensure_ascii=False, indent=2)
            
except KeyboardInterrupt:
    print("\n중단됨. 캐시 저장 완료.")

found_count = sum(1 for v in osm_cache.values() if v.get('lat'))
print(f"\n✨ 완료! 총 OSM 확보 역 수: {found_count}개")
